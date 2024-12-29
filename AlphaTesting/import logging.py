import logging
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from flask_pymongo import PyMongo
from werkzeug.security import generate_password_hash, check_password_hash
from bson.objectid import ObjectId
import jwt
from functools import wraps
from datetime import datetime, timedelta
import os
import torch
from transformers import RobertaTokenizer, RobertaForSequenceClassification, RobertaModel
from PyPDF2 import PdfReader
from docx import Document
import tempfile
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from nudenet import NudeDetector
from flasgger import Swagger, swag_from
from pymongo.errors import ConnectionFailure
from reportlab.lib import colors
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.piecharts import Pie
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.piecharts import Pie
from reportlab.lib.units import inch
import cv2
import numpy as np
from moviepy.editor import VideoFileClip, CompositeVideoClip, TextClip

import tempfile
import os
import cv2
import numpy as np
from moviepy.editor import VideoFileClip, CompositeVideoClip, TextClip
from nudenet import NudeDetector
import tempfile
import os
import shutil
import logging


logging.basicConfig(level=logging.DEBUG, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "*"}})
Swagger(app)


app.config['SECRET_KEY'] = 'your_secret_key'
app.config['MONGO_URI'] = 'mongodb://localhost:27017/document_analyzer'
mongo = PyMongo(app)

class RobertaClass(torch.nn.Module):
    def __init__(self):
        super(RobertaClass, self).__init__()
        self.l1 = RobertaModel.from_pretrained("roberta-base")
        self.pre_classifier = torch.nn.Linear(768, 768)
        self.dropout = torch.nn.Dropout(0.3)
        self.classifier = torch.nn.Linear(768, 2)  

    def forward(self, input_ids, attention_mask):
        outputs = self.l1(input_ids=input_ids, attention_mask=attention_mask)
        hidden_state = outputs[0]
        pooler = hidden_state[:, 0]
        pooler = self.pre_classifier(pooler)
        pooler = torch.nn.ReLU()(pooler)
        pooler = self.dropout(pooler)
        output = self.classifier(pooler)
        return output
    

logger.info("Initializing RoBERTa model...")
model = RobertaClass()
tokenizer = RobertaTokenizer.from_pretrained(r'./tokenizer')
try:
    model.load_state_dict(torch.load("pytorch_roberta_cyberbullying.bin", map_location=torch.device('cpu')))
    model.to('cpu')
    model.eval()
    logger.info("RoBERTa model loaded successfully")
except Exception as e:
    logger.error(f"Error loading RoBERTa model: {str(e)}")
    


logger.info("Initializing NudeDetector...")
nude_detector = NudeDetector()
logger.info("NudeDetector initialized")


try:
    
    mongo.db.command('ismaster')
    logger.info("MongoDB connection successful")
except ConnectionFailure:
    logger.error("MongoDB server not available")
    

def token_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        token = request.headers.get('Authorization')
        if not token:
            logger.warning("No token provided in request")
            return jsonify({'message': 'Token is missing!'}), 401
        try:
            data = jwt.decode(token, app.config['SECRET_KEY'], algorithms=["HS256"])
            current_user = mongo.db.users.find_one({'_id': ObjectId(data['user_id'])})
            logger.info(f"User authenticated: {current_user['email']}")
        except Exception as e:
            logger.error(f"Token validation failed: {str(e)}")
            return jsonify({'message': 'Token is invalid!'}), 401
        return f(current_user, *args, **kwargs)
    return decorated

@app.route('/register', methods=['POST'])
@swag_from({
    'tags': ['Authentication'],
    'description': 'Register a new user',
    'parameters': [
        {
            'name': 'body',
            'in': 'body',
            'schema': {
                'type': 'object',
                'properties': {
                    'email': {'type': 'string'},
                    'password': {'type': 'string'}
                }
            }
        }
    ],
    'responses': {
        '201': {
            'description': 'User created successfully'
        },
        '400': {
            'description': 'User already exists'
        }
    }
})
def register():
    data = request.get_json()
    logger.info(f"Attempting to register user: {data['email']}")
    hashed_password = generate_password_hash(data['password'])
    new_user = {
        'email': data['email'],
        'password': hashed_password
    }
    if mongo.db.users.find_one({'email': data['email']}):
        logger.warning(f"Registration failed: User {data['email']} already exists")
        return jsonify({'message': 'User already exists'}), 400
    mongo.db.users.insert_one(new_user)
    logger.info(f"User registered successfully: {data['email']}")
    return jsonify({'message': 'User created successfully'}), 201

@app.route('/login', methods=['POST'])
@swag_from({
    'tags': ['Authentication'],
    'description': 'Login and receive a JWT token',
    'parameters': [
        {
            'name': 'body',
            'in': 'body',
            'schema': {
                'type': 'object',
                'properties': {
                    'email': {'type': 'string'},
                    'password': {'type': 'string'}
                }
            }
        }
    ],
    'responses': {
        '200': {
            'description': 'Login successful',
            'schema': {
                'type': 'object',
                'properties': {
                    'token': {'type': 'string'}
                }
            }
        },
        '401': {
            'description': 'Invalid credentials'
        }
    }
})
def login():
    data = request.get_json()
    logger.info(f"Login attempt for user: {data['email']}")
    user = mongo.db.users.find_one({'email': data['email']})
    if user and check_password_hash(user['password'], data['password']):
        token = jwt.encode({
            'user_id': str(user['_id']),
            'exp': datetime.utcnow() + timedelta(hours=24)
        }, app.config['SECRET_KEY'])
        logger.info(f"Login successful for user: {data['email']}")
        return jsonify({'token': token})
    logger.warning(f"Login failed for user: {data['email']}")
    return jsonify({'message': 'Invalid credentials'}), 401

def analyze_text(text):
    logger.info("Starting text analysis")
    words = text.split()
    word_scores = []
    
    for word in words:
        inputs = tokenizer(word, return_tensors="pt", truncation=True, max_length=512, padding=True)
        with torch.no_grad():
            outputs = model(input_ids=inputs['input_ids'], 
                            attention_mask=inputs['attention_mask'])
        probabilities = torch.nn.functional.softmax(outputs, dim=-1)
        word_scores.append(probabilities.tolist()[0][1])  
    
    overall_score = sum(word_scores) / len(word_scores)
    logger.info(f"Text analysis complete. Overall score: {overall_score}")
    return overall_score, list(zip(words, word_scores))


def analyze_image(image_path):
    logger.info(f"Starting image analysis for: {image_path}")
    result = nude_detector.detect(image_path)
    logger.debug(f"NudeDetector result: {result}")
    score = sum(detection['score'] for detection in result if detection['score'] > 0.5) / len(result) if result else 0
    logger.info(f"Image analysis complete. Score: {score}")
    return score

def extract_text(file):
    logger.info(f"Extracting text from file: {file.filename}")
    try:
        _, ext = os.path.splitext(file.filename)
        if ext.lower() == '.pdf':
            logger.debug("Extracting text from PDF")
            reader = PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
        elif ext.lower() in ['.doc', '.docx']:
            logger.debug("Extracting text from Word document")
            doc = Document(file)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext.lower() == '.txt':
            logger.debug("Extracting text from TXT file")
            text = file.read().decode('utf-8')
        else:
            logger.error(f"Unsupported file format: {ext}")
            raise ValueError("Unsupported file format")
        logger.info(f"Text extraction complete. Extracted {len(text)} characters")
        return text
    except Exception as e:
        logger.error(f"Error extracting text: {str(e)}")
        raise

from reportlab.lib import colors
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.piecharts import Pie

def generate_report(text_score, word_scores, image_score=None):
    logger.info("Generating analysis report")
    report_path = tempfile.mktemp(suffix='.pdf')
    doc = SimpleDocTemplate(report_path, pagesize=letter,
                            rightMargin=72, leftMargin=72,
                            topMargin=72, bottomMargin=18)

    Story = []
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Justify', alignment=4))  

    
    header_style = styles['Heading1']
    header_style.alignment = 1  
    Story.append(Paragraph("Verified by NoHate", header_style))
    Story.append(Spacer(1, 12))

    
    Story.append(Paragraph(f"Overall Text Content Score: {text_score:.2f}", styles['Normal']))
    if image_score is not None:
        Story.append(Paragraph(f"Image Content Score: {image_score:.2f}", styles['Normal']))
    Story.append(Spacer(1, 12))

    
    text = []
    for word, score in word_scores:
        color = f"
        text.append(f'<font color="{color}">{word}</font>')
    
    highlighted_text = ' '.join(text)
    Story.append(Paragraph(highlighted_text, styles['Justify']))
    Story.append(Spacer(1, 24))

    
    d = Drawing(400, 200)
    pc = Pie()
    pc.x = 65
    pc.y = 15
    pc.width = 300
    pc.height = 150
    pc.data = [text_score, 1 - text_score]
    pc.labels = ['Problematic', 'Safe']
    pc.slices.strokeWidth = 0.5
    pc.slices[0].fillColor = colors.red
    pc.slices[1].fillColor = colors.green
    d.add(pc)
    
    Story.append(Paragraph("Content Analysis Statistics", styles['Heading2']))
    Story.append(d)

    doc.build(Story)
    
    logger.info(f"Report generated successfully: {report_path}")
    return report_path

@app.route('/analyze', methods=['POST'])
@token_required
@swag_from({
    'tags': ['Document Analysis'],
    'description': 'Analyze a document, image, or video for inappropriate content',
    'parameters': [
        {
            'name': 'Authorization',
            'in': 'header',
            'type': 'string',
            'required': 'true',
            'description': 'JWT token'
        },
        {
            'name': 'file',
            'in': 'formData',
            'type': 'file',
            'required': 'true',
            'description': 'Document, image, or video to analyze'
        }
    ],
    'responses': {
        '200': {
            'description': 'Analysis report, image stats, or censored video',
            'content': {
                'application/pdf': {},
                'application/json': {},
                'video/mp4': {}
            }
        },
        '400': {
            'description': 'Bad request'
        },
        '401': {
            'description': 'Unauthorized'
        },
        '500': {
            'description': 'Internal server error'
        }
    }
})
def analyze_document(current_user):
    logger.info(f"Analysis request received from user: {current_user['email']}")
    if 'file' not in request.files:
        logger.warning("No file part in the request")
        return jsonify({'message': 'No file part in the request'}), 400
    file = request.files['file']
    if file.filename == '':
        logger.warning("No file selected for uploading")
        return jsonify({'message': 'No file selected for uploading'}), 400
    
    logger.info(f"Processing file: {file.filename}")
    try:
        _, ext = os.path.splitext(file.filename)
        if ext.lower() in ['.pdf', '.doc', '.docx', '.txt']:
            logger.info("Extracting text from document")
            text = extract_text(file)
            logger.info("Analyzing extracted text")
            text_score, word_scores = analyze_text(text)
            logger.info("Generating report for text analysis")
            report_path = generate_report(text_score, word_scores)
            return send_file(report_path, as_attachment=True, download_name='analysis_report.pdf', mimetype='application/pdf')
        elif ext.lower() in ['.jpg', '.jpeg', '.png']:
            logger.info("Processing image file")
            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
                file.save(temp_file.name)
                logger.info("Analyzing image")
                image_stats = analyze_image(temp_file.name)
            os.unlink(temp_file.name)
            return jsonify(image_stats), 200
        elif ext.lower() in ['.mp4', '.avi', '.mov']:
            logger.info("Processing video file")
            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
                file.save(temp_file.name)
                logger.info("Censoring video")
                censored_video_path = process_video(temp_file.name)
            logger.info("Sending censored video")
            return send_file(censored_video_path, as_attachment=True, download_name='censored_video.mp4', mimetype='video/mp4')
        else:
            logger.error(f"Unsupported file format: {ext}")
            return jsonify({'message': 'Unsupported file format'}), 400
        
    except Exception as e:
        logger.error(f"Error during analysis: {str(e)}", exc_info=True)
        return jsonify({'message': 'An error occurred during analysis'}), 500




def process_video(input_path):
    logger.info(f"Processing video: {input_path}")
    detector = NudeDetector()
    
    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            
            cap = cv2.VideoCapture(input_path)
            if not cap.isOpened():
                logger.error("Error opening video file")
                raise ValueError("Could not open video file")

            
            fps = cap.get(cv2.CAP_PROP_FPS)
            width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
            height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            
            
            fourcc = cv2.VideoWriter_fourcc(*'mp4v')
            out_path = os.path.join(temp_dir, 'censored_video.mp4')
            out = cv2.VideoWriter(out_path, fourcc, fps, (width, height))
            
            frame_count = 0
            while True:
                ret, frame = cap.read()
                if not ret:
                    break
                
                logger.debug(f"Processing frame {frame_count}/{total_frames}")
                
                try:
                    
                    detections = detector.detect(frame)
                    for detection in detections:
                        if detection['score'] > 0.5:  
                            box = detection['box']
                            x1, y1, x2, y2 = map(int, box)
                            
                            
                            x1 = max(0, min(x1, width - 1))
                            y1 = max(0, min(y1, height - 1))
                            x2 = max(x1 + 1, min(x2, width))
                            y2 = max(y1 + 1, min(y2, height))
                            
                            
                            frame[y1:y2, x1:x2] = cv2.blur(frame[y1:y2, x1:x2], (30, 30))
                            
                            logger.info(f"Censored area: [{x1}, {y1}, {x2}, {y2}]")
                except Exception as e:
                    logger.error(f"Error processing frame {frame_count}: {str(e)}")
                    
                    continue
                
                
                out.write(frame)
                frame_count += 1
                
                if frame_count % 100 == 0:
                    logger.info(f"Processed {frame_count}/{total_frames} frames")
            
            cap.release()
            out.release()
            
            logger.info("Adding watermark to video")
            
            video = VideoFileClip(out_path)
            watermark = (TextClip("Verified by NoHate", fontsize=24, color='white')
                         .set_position(('right', 'bottom'))
                         .set_duration(video.duration))
            
            final_video = CompositeVideoClip([video, watermark])
            final_output_path = os.path.join(temp_dir, 'final_censored_video.mp4')
            final_video.write_videofile(final_output_path, codec='libx264')
            
            
            permanent_output_path = os.path.join(os.path.dirname(input_path), 'censored_' + os.path.basename(input_path))
            shutil.copy2(final_output_path, permanent_output_path)
            
            logger.info(f"Video processing complete. Output saved to: {permanent_output_path}")
            return permanent_output_paqth
        
        except Exception as e:
            logger.error(f"Error processing video: {str(e)}", exc_info=True)
            raise
        finally:
            if 'cap' in locals() and cap.isOpened():
                cap.release()
            if 'out' in locals():
                out.release()

import cv2
import numpy as np
from moviepy.editor import VideoFileClip, CompositeVideoClip, TextClip
from nudenet import NudeDetector
import tempfile
import os
import shutil
import logging

import moviepy.config as cfg
cfg.change_settings({"IMAGEMAGICK_BINARY": "magick.exe"})

def analyze_image(image_path):
    logger.info(f"Starting image analysis for: {image_path}")
    detector = NudeDetector()
    result = detector.detect(image_path)
    logger.debug(f"NudeDetector result: {result}")

    
    overall_score = sum(detection['score'] for detection in result) / len(result) if result else 0

    
    part_scores = {}
    for detection in result:
        part = detection['class']
        score = detection['score']
        if part not in part_scores:
            part_scores[part] = []
        part_scores[part].append(score)

    
    avg_part_scores = {part: sum(scores) / len(scores) for part, scores in part_scores.items()}

    
    part_counts = {part: len(scores) for part, scores in part_scores.items()}

    
    stats = {
        'overall_score': overall_score,
        'total_detections': len(result),
        'part_scores': avg_part_scores,
        'part_counts': part_counts,
    }

    logger.info(f"Image analysis complete. Overall score: {overall_score}")
    logger.debug(f"Detailed stats: {stats}")

    return stats

if __name__ == '__main__':
    logger.info("Starting Flask application")
    app.run(debug=True)