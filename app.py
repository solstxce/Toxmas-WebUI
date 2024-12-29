import logging
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from flask_pymongo import PyMongo
from werkzeug.security import generate_password_hash, check_password_hash
from bson.objectid import ObjectId
import jwt
from functools import wraps
from datetime import datetime, timedelta
import os
import torch
from transformers import RobertaTokenizer, RobertaForSequenceClassification, RobertaModel
from PyPDF2 import PdfReader
from docx import Document
import tempfile
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from nudenet import NudeDetector
from flasgger import Swagger, swag_from
from pymongo.errors import ConnectionFailure
from reportlab.lib import colors
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.piecharts import Pie
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.piecharts import Pie
from reportlab.lib.units import inch
import cv2
import numpy as np
from moviepy.editor import VideoFileClip, CompositeVideoClip, TextClip
# from nudenet import NudeDetector
import tempfile
import os
import cv2
import numpy as np
from moviepy.editor import VideoFileClip, CompositeVideoClip, TextClip
from nudenet import NudeDetector
import tempfile
import os
import shutil
import logging
from better_profanity import profanity
import nltk
from nltk.corpus import wordnet
import re
from moviepy.editor import VideoFileClip, concatenate_videoclips, AudioFileClip
import whisper
from pydub import AudioSegment

# Configure logging
logging.basicConfig(level=logging.DEBUG, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "*"}})
Swagger(app)

# Configuration
app.config['SECRET_KEY'] = 'your_secret_key'
app.config['MONGO_URI'] = 'mongodb://localhost:27017/document_analyzer'
mongo = PyMongo(app)

class RobertaClass(torch.nn.Module):
    def __init__(self):
        super(RobertaClass, self).__init__()
        self.l1 = RobertaModel.from_pretrained("roberta-base")
        self.pre_classifier = torch.nn.Linear(768, 768)
        self.dropout = torch.nn.Dropout(0.3)
        self.classifier = torch.nn.Linear(768, 2)  # Binary classification

    def forward(self, input_ids, attention_mask):
        outputs = self.l1(input_ids=input_ids, attention_mask=attention_mask)
        hidden_state = outputs[0]
        pooler = hidden_state[:, 0]
        pooler = self.pre_classifier(pooler)
        pooler = torch.nn.ReLU()(pooler)
        pooler = self.dropout(pooler)
        output = self.classifier(pooler)
        return output
    
# Custom RoBERTa Model Setup
logger.info("Initializing RoBERTa model...")
model = RobertaClass()
tokenizer = RobertaTokenizer.from_pretrained(r'./tokenizer')
try:
    model.load_state_dict(torch.load("pytorch_roberta_cyberbullying.bin", map_location=torch.device('cpu')))
    model.to('cpu')
    model.eval()
    logger.info("RoBERTa model loaded successfully")
except Exception as e:
    logger.error(f"Error loading RoBERTa model: {str(e)}")
    # Handle the error appropriately, maybe set a flag to disable analysis if model can't be loaded

# NudeNet Setup
logger.info("Initializing NudeDetector...")
nude_detector = NudeDetector()
logger.info("NudeDetector initialized")

# Verify MongoDB connection
try:
    # The ismaster command is cheap and does not require auth.
    mongo.db.command('ismaster')
    logger.info("MongoDB connection successful")
except ConnectionFailure:
    logger.error("MongoDB server not available")
    # Handle the error appropriately

def token_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        token = request.headers.get('Authorization')
        if not token:
            logger.warning("No token provided in request")
            return jsonify({'message': 'Token is missing!'}), 401
        try:
            data = jwt.decode(token, app.config['SECRET_KEY'], algorithms=["HS256"])
            current_user = mongo.db.users.find_one({'_id': ObjectId(data['user_id'])})
            logger.info(f"User authenticated: {current_user['email']}")
        except Exception as e:
            logger.error(f"Token validation failed: {str(e)}")
            return jsonify({'message': 'Token is invalid!'}), 401
        return f(current_user, *args, **kwargs)
    return decorated

@app.route('/register', methods=['POST'])
@swag_from({
    'tags': ['Authentication'],
    'description': 'Register a new user',
    'parameters': [
        {
            'name': 'body',
            'in': 'body',
            'schema': {
                'type': 'object',
                'properties': {
                    'email': {'type': 'string'},
                    'password': {'type': 'string'}
                }
            }
        }
    ],
    'responses': {
        '201': {
            'description': 'User created successfully'
        },
        '400': {
            'description': 'User already exists'
        }
    }
})
def register():
    data = request.get_json()
    logger.info(f"Attempting to register user: {data['email']}")
    hashed_password = generate_password_hash(data['password'])
    new_user = {
        'email': data['email'],
        'password': hashed_password
    }
    if mongo.db.users.find_one({'email': data['email']}):
        logger.warning(f"Registration failed: User {data['email']} already exists")
        return jsonify({'message': 'User already exists'}), 400
    mongo.db.users.insert_one(new_user)
    logger.info(f"User registered successfully: {data['email']}")
    return jsonify({'message': 'User created successfully'}), 201

@app.route('/login', methods=['POST'])
@swag_from({
    'tags': ['Authentication'],
    'description': 'Login and receive a JWT token',
    'parameters': [
        {
            'name': 'body',
            'in': 'body',
            'schema': {
                'type': 'object',
                'properties': {
                    'email': {'type': 'string'},
                    'password': {'type': 'string'}
                }
            }
        }
    ],
    'responses': {
        '200': {
            'description': 'Login successful',
            'schema': {
                'type': 'object',
                'properties': {
                    'token': {'type': 'string'}
                }
            }
        },
        '401': {
            'description': 'Invalid credentials'
        }
    }
})
def login():
    data = request.get_json()
    logger.info(f"Login attempt for user: {data['email']}")
    user = mongo.db.users.find_one({'email': data['email']})
    if user and check_password_hash(user['password'], data['password']):
        token = jwt.encode({
            'user_id': str(user['_id']),
            'exp': datetime.utcnow() + timedelta(hours=24)
        }, app.config['SECRET_KEY'])
        logger.info(f"Login successful for user: {data['email']}")
        return jsonify({'token': token})
    logger.warning(f"Login failed for user: {data['email']}")
    return jsonify({'message': 'Invalid credentials'}), 401

def get_legal_alternatives(word):
    alternatives = []
    for syn in wordnet.synsets(word):
        for lemma in syn.lemmas():
            if lemma.name() != word and not profanity.contains_profanity(lemma.name()):
                alternatives.append(lemma.name().replace('_', ' '))
    return alternatives if alternatives else [word]

def replace_profanity_in_sentence(sentence):
    words = sentence.split()
    replaced_sentence = []

    for word in words:
        cleaned_word = re.sub(r'[^\w\s]', '', word)
        if profanity.contains_profanity(cleaned_word):
            alternatives = get_legal_alternatives(cleaned_word)
            replaced_sentence.append(alternatives[0])
        else:
            replaced_sentence.append(word)

    return ' '.join(replaced_sentence)

def replace_profanity_in_text(text):
    sentences = re.split(r'(?<=[.!?]) +', text)
    replaced_sentences = [replace_profanity_in_sentence(sentence) for sentence in sentences]
    return ' '.join(replaced_sentences)

def check_and_replace_profanity(text):
    contains_profanity = profanity.contains_profanity(text)
    replaced_text = replace_profanity_in_text(text)
    return contains_profanity, replaced_text

def analyze_text(text, mode='censor'):
    logger.info(f"Starting text analysis in {mode} mode")
    words = text.split()
    word_scores = []
    
    for word in words:
        inputs = tokenizer(word, return_tensors="pt", truncation=True, max_length=512, padding=True)
        with torch.no_grad():
            outputs = model(input_ids=inputs['input_ids'], 
                          attention_mask=inputs['attention_mask'])
        probabilities = torch.nn.functional.softmax(outputs, dim=-1)
        word_scores.append(probabilities.tolist()[0][1])
    
    overall_score = sum(word_scores) / len(word_scores)
    
    # Apply censoring or rephrasing based on mode
    if mode == 'censor':
        processed_text = profanity.censor(text)
    else:  # rephrase mode
        contains_profanity, processed_text = check_and_replace_profanity(text)
    
    logger.info(f"Text analysis complete. Overall score: {overall_score}")
    return overall_score, list(zip(words, word_scores)), processed_text

def analyze_image(image_path):
    logger.info(f"Starting image analysis for: {image_path}")
    result = nude_detector.detect(image_path)
    logger.debug(f"NudeDetector result: {result}")
    score = sum(detection['score'] for detection in result if detection['score'] > 0.5) / len(result) if result else 0
    logger.info(f"Image analysis complete. Score: {score}")
    return score

def extract_text(file):
    logger.info(f"Extracting text from file: {file.filename}")
    try:
        _, ext = os.path.splitext(file.filename)
        if ext.lower() == '.pdf':
            logger.debug("Extracting text from PDF")
            reader = PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
        elif ext.lower() in ['.doc', '.docx']:
            logger.debug("Extracting text from Word document")
            doc = Document(file)
            text = "\n".join([para.text for para in doc.paragraphs])
        elif ext.lower() == '.txt':
            logger.debug("Extracting text from TXT file")
            text = file.read().decode('utf-8')
        else:
            logger.error(f"Unsupported file format: {ext}")
            raise ValueError("Unsupported file format")
        logger.info(f"Text extraction complete. Extracted {len(text)} characters")
        return text
    except Exception as e:
        logger.error(f"Error extracting text: {str(e)}")
        raise


from reportlab.lib import colors
from reportlab.graphics.shapes import Drawing
from reportlab.graphics.charts.piecharts import Pie

def generate_report(text_score, word_scores, processed_text, image_score=None):
    logger.info("Generating analysis report")
    report_path = tempfile.mktemp(suffix='.pdf')
    doc = SimpleDocTemplate(report_path, pagesize=letter,
                            rightMargin=72, leftMargin=72,
                            topMargin=72, bottomMargin=18)

    Story = []
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Justify', alignment=4))  # 4 is for full justification

    # Add header
    header_style = styles['Heading1']
    header_style.alignment = 1  # Center alignment
    Story.append(Paragraph("Verified by Tox-Mas", header_style))
    Story.append(Spacer(1, 12))

    # Add overall scores
    Story.append(Paragraph(f"Overall Text Content Score: {text_score:.2f}", styles['Normal']))
    if image_score is not None:
        Story.append(Paragraph(f"Image Content Score: {image_score:.2f}", styles['Normal']))
    Story.append(Spacer(1, 12))

    # Add processed text
    Story.append(Paragraph("Processed Text:", styles['Heading2']))
    Story.append(Paragraph(processed_text, styles['Justify']))
    Story.append(Spacer(1, 24))

    # Word-level highlighting
    text = []
    for word, score in word_scores:
        color = f"#{int(min(1, score * 2) * 255):02x}{int(max(0, 1 - score * 2) * 255):02x}00"
        text.append(f'<font color="{color}">{word}</font>')
    
    highlighted_text = ' '.join(text)
    Story.append(Paragraph(highlighted_text, styles['Justify']))
    Story.append(Spacer(1, 24))

    # Pie chart for statistics
    d = Drawing(400, 200)
    pc = Pie()
    pc.x = 65
    pc.y = 15
    pc.width = 300
    pc.height = 150
    pc.data = [text_score, 1 - text_score]
    pc.labels = ['Problematic', 'Safe']
    pc.slices.strokeWidth = 0.5
    pc.slices[0].fillColor = colors.red
    pc.slices[1].fillColor = colors.green
    d.add(pc)
    
    Story.append(Paragraph("Content Analysis Statistics", styles['Heading2']))
    Story.append(d)

    doc.build(Story)
    
    logger.info(f"Report generated successfully: {report_path}")
    return report_path


@app.route('/analyze', methods=['POST'])
@token_required
@swag_from({
    'tags': ['Document Analysis'],
    'description': 'Analyze a document, image, or video for inappropriate content',
    'parameters': [
        {
            'name': 'Authorization',
            'in': 'header',
            'type': 'string',
            'required': 'true',
            'description': 'JWT token'
        },
        {
            'name': 'file',
            'in': 'formData',
            'type': 'file',
            'required': 'true',
            'description': 'Document, image, or video to analyze'
        }
    ],
    'responses': {
        '200': {
            'description': 'Analysis report, image stats, or censored video',
            'content': {
                'application/pdf': {},
                'application/json': {},
                'video/mp4': {},
                'image/jpeg': {},
                'audio/wav': {},
                'audio/mp3': {}
            }
        },
        '400': {
            'description': 'Bad request'
        },
        '401': {
            'description': 'Unauthorized'
        },
        '500': {
            'description': 'Internal server error'
        }
    }
})
def analyze_document(current_user):
    logger.info(f"Analysis request received from user: {current_user['email']}")
    if 'file' not in request.files:
        logger.warning("No file part in the request")
        return jsonify({'message': 'No file part in the request'}), 400
    file = request.files['file']
    if file.filename == '':
        logger.warning("No file selected for uploading")
        return jsonify({'message': 'No file selected for uploading'}), 400
    
    logger.info(f"Processing file: {file.filename}")
    try:
        _, ext = os.path.splitext(file.filename)
        if ext.lower() in ['.pdf', '.doc', '.docx', '.txt']:
            text = extract_text(file)
            mode = request.form.get('mode', 'censor')  # Default to censor if not specified
            text_score, word_scores, processed_text = analyze_text(text, mode)
            report_path = generate_report(text_score, word_scores, processed_text)
            return send_file(report_path, as_attachment=True, download_name='analysis_report.pdf', mimetype='application/pdf')
        elif ext.lower() in ['.jpg', '.jpeg', '.png']:
            logger.info("Processing image file")
            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
                file.save(temp_file.name)
                logger.info("Analyzing image")
                image_stats = analyze_image(temp_file.name)
            os.unlink(temp_file.name)
            return jsonify(image_stats), 200
        elif ext.lower() in ['.mp4', '.avi', '.mov']:
            logger.info("Processing video file")
            with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as temp_file:
                file.save(temp_file.name)
                logger.info("Censoring video")
                censored_video_path = process_video(temp_file.name)
            logger.info("Sending censored video")
            return send_file(censored_video_path, as_attachment=True, download_name='censored_video.mp4', mimetype='video/mp4')
        else:
            logger.error(f"Unsupported file format: {ext}")
            return jsonify({'message': 'Unsupported file format'}), 400
        
    except Exception as e:
        logger.error(f"Error during analysis: {str(e)}", exc_info=True)
        return jsonify({'message': 'An error occurred during analysis'}), 500


import cv2
import numpy as np
from moviepy.editor import VideoFileClip, CompositeVideoClip, TextClip
from nudenet import NudeDetector
import tempfile
import os
import shutil
import logging
from PIL import Image
import matplotlib.pyplot as plt
import math

logger = logging.getLogger(__name__)

def process_video(input_path):
    logger.info(f"Processing video: {input_path}")
    detector = NudeDetector(classes= [
    "FEMALE_GENITALIA_COVERED",
    "BUTTOCKS_EXPOSED",
    "FEMALE_BREAST_EXPOSED",
    "FEMALE_GENITALIA_EXPOSED",
    "ANUS_EXPOSED",
    "BELLY_COVERED",
    "ARMPITS_COVERED",
    "ARMPITS_EXPOSED",
    "BELLY_EXPOSED",
    "MALE_GENITALIA_EXPOSED",
    "ANUS_COVERED",
    "FEMALE_BREAST_COVERED",
    "BUTTOCKS_COVERED",
])
    
    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            cap = cv2.VideoCapture(input_path)
            if not cap.isOpened():
                logger.error("Error opening video file")
                raise ValueError("Could not open video file")

            fps = cap.get(cv2.CAP_PROP_FPS)
            width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
            height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            
            fourcc = cv2.VideoWriter_fourcc(*'mp4v')
            out_path = os.path.join(temp_dir, 'censored_video.mp4')
            out = cv2.VideoWriter(out_path, fourcc, fps, (width, height))
            
            frame_count = 0
            while True:
                ret, frame = cap.read()
                if not ret:
                    break
                
                if frame_count % 5 == 0:  # Process every 5th frame for performance
                    logger.debug(f"Processing frame {frame_count}/{total_frames}")
                    
                    try:
                        # Use detector.censor instead of detect
                        censored_frame = detector.censor(frame)
                        if censored_frame is not None:
                            frame = censored_frame
                        else:
                            logger.warning(f"Censorship returned None for frame {frame_count}")
                    except Exception as e:
                        logger.error(f"Error processing frame {frame_count}: {str(e)}")
                    # Continue processing other frames even if one fails
                
                out.write(frame)
                frame_count += 1
                
                if frame_count % 100 == 0:
                    logger.info(f"Processed {frame_count}/{total_frames} frames")
            
            cap.release()
            out.release()
            
            logger.info("Adding watermark to video")
            
            video = VideoFileClip(out_path)
            watermark = (TextClip("Verified by NoHate", fontsize=24, color='white', font='Arial')
                         .set_position(('right', 'bottom'))
                         .set_duration(video.duration))
            
            final_video = CompositeVideoClip([video, watermark])
            final_output_path = os.path.join(temp_dir, 'final_censored_video.mp4')
            final_video.write_videofile(final_output_path, codec='libx264', audio_codec='aac')
            
            permanent_output_path = os.path.join(os.path.dirname(input_path), 'censored_' + os.path.basename(input_path))
            shutil.copy2(final_output_path, permanent_output_path)
            
            logger.info(f"Video processing complete. Output saved to: {permanent_output_path}")
            return permanent_output_path
        
        except Exception as e:
            logger.error(f"Error processing video: {str(e)}", exc_info=True)
            raise
        finally:
            if 'cap' in locals() and cap.isOpened():
                cap.release()
            if 'out' in locals():
                out.release()

def analyze_image(image_path):
    logger.info(f"Starting image analysis for: {image_path}")
    detector = NudeDetector()
    
    try:
        # Use detector.censor instead of detect
        censored_img_path = detector.censor(image_path)
        if censored_img_path:
            censored_img = Image.open(censored_img_path)
            
            # Calculate overall score based on the difference between original and censored image
            original_img = Image.open(image_path)
            diff = np.array(original_img) - np.array(censored_img)
            overall_score = np.sum(np.abs(diff)) / (original_img.size[0] * original_img.size[1] * 3)  # Normalize by image size and channels
            
            stats = {
                'overall_score': overall_score,
                'censored_image_path': censored_img_path
            }
            
            logger.info(f"Image analysis complete. Overall score: {overall_score}")
            logger.debug(f"Detailed stats: {stats}")
            
            return stats
        else:
            logger.warning(f"Censorship returned None for {image_path}")
            return None
    except Exception as e:
        logger.error(f"Error analyzing image {image_path}: {str(e)}")
        return None


import cv2
import numpy as np
from moviepy.editor import VideoFileClip, CompositeVideoClip, TextClip
from nudenet import NudeDetector
import tempfile
import os
import shutil
import logging

import moviepy.config as cfg
cfg.change_settings({"IMAGEMAGICK_BINARY": "magick.exe"})

def extract_audio(video_path):
    """Extract audio from video file."""
    video = VideoFileClip(video_path)
    audio_path = "extracted_audio.wav"
    video.audio.write_audiofile(audio_path)
    video.close()
    return audio_path

def transcribe_audio_whisper(audio_path):
    """Transcribe audio using Whisper model."""
    model = whisper.load_model("base")
    result = model.transcribe(audio_path)
    return result["text"]

def convert_audio_to_text_with_timestamps_whisper(audio_path, transcribed_text):
    """Use Whisper-transcribed text and assign timestamps."""
    audio = AudioSegment.from_wav(audio_path)
    chunk_length_ms = 2000
    chunks = [audio[i:i + chunk_length_ms] for i in range(0, len(audio), chunk_length_ms)]

    transcription = []
    words = transcribed_text.split()
    chunk_size = len(words) // len(chunks) if len(chunks) > 0 else len(words)

    for i, chunk in enumerate(chunks):
        start_time = i * 2
        chunk_text = " ".join(words[i * chunk_size:(i + 1) * chunk_size])
        transcription.append((start_time, chunk_text))

    return transcription

def timestamp_to_seconds(timestamp):
    minutes, seconds = map(int, timestamp.split(':'))
    return minutes * 60 + seconds

def find_harmful_word_timestamps(transcript_lines, timestamps):
    mute_times = []
    for i, line in enumerate(transcript_lines):
        if profanity.contains_profanity(line):
            mute_times.append(timestamps[i])
    return mute_times

def mute_video_at_timestamps(video_file, mute_times):
    video = VideoFileClip(video_file)
    clips = []
    last_end = 0

    for timestamp in mute_times:
        start_time = timestamp_to_seconds(timestamp)
        end_time = start_time + 1

        if start_time > last_end:
            clips.append(video.subclip(last_end, start_time))

        muted_clip = video.subclip(start_time, end_time).volumex(0)
        clips.append(muted_clip)
        last_end = end_time

    if last_end < video.duration:
        clips.append(video.subclip(last_end, video.duration))

    final_clip = concatenate_videoclips(clips)
    output_path = os.path.join(os.path.dirname(video_file), 'audio_censored_' + os.path.basename(video_file))
    final_clip.write_videofile(output_path, codec="libx264", audio_codec="aac")
    return output_path

def process_audio(file_path):
    """Process audio file or video's audio for censoring."""
    logger.info("Processing audio for censoring")
    
    try:
        # Extract audio if it's a video file
        _, ext = os.path.splitext(file_path)
        if ext.lower() in ['.mp4', '.avi', '.mov']:
            audio_path = extract_audio(file_path)
            is_video = True
        else:
            audio_path = file_path
            is_video = False

        # Transcribe and process
        transcribed_text = transcribe_audio_whisper(audio_path)
        transcription = convert_audio_to_text_with_timestamps_whisper(audio_path, transcribed_text)
        
        timestamps = [f"{int(ts // 60):02d}:{int(ts % 60):02d}" for ts, _ in transcription]
        transcript_lines = [text for _, text in transcription]
        mute_times = find_harmful_word_timestamps(transcript_lines, timestamps)

        if is_video:
            # If input was video, mute the video at specified timestamps
            output_path = mute_video_at_timestamps(file_path, mute_times)
            os.remove(audio_path)  # Clean up extracted audio
        else:
            # If input was audio, create muted audio file
            audio = AudioSegment.from_wav(audio_path)
            for timestamp in mute_times:
                start_time = timestamp_to_seconds(timestamp) * 1000  # Convert to milliseconds
                end_time = start_time + 1000  # Mute for 1 second
                muted_segment = AudioSegment.silent(duration=end_time-start_time)
                audio = audio[:start_time] + muted_segment + audio[end_time:]
            
            output_path = os.path.join(os.path.dirname(file_path), 'censored_' + os.path.basename(file_path))
            audio.export(output_path, format="wav")

        logger.info(f"Audio processing complete. Output saved to: {output_path}")
        return output_path

    except Exception as e:
        logger.error(f"Error processing audio: {str(e)}", exc_info=True)
        raise

if __name__ == '__main__':
    logger.info("Starting Flask application")
    app.run(debug=True)